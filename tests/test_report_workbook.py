from pathlib import Path

import openpyxl

from fncsqaqc.models import CheckResult, ConversionLogEntry, RunResult, Severity
from fncsqaqc.report.workbook_builder import build


def test_build_report_smoke(tmp_path):
    run = RunResult(folder=Path("D:/Projects/Test"))
    run.add(CheckResult(check_id="layer_names", severity=Severity.FAIL, file="D-101.dwg", message="bad layer"))
    run.add(CheckResult(check_id="purge", severity=Severity.WARN, file="D-101.dwg", message="unused layer"))
    run.conversion_log.append(ConversionLogEntry("D-101.dwg", "converted", duration_s=1.2))

    output = tmp_path / "report.xlsx"
    build(run, output)

    assert output.is_file()
    wb = openpyxl.load_workbook(output)
    expected_sheets = {
        "Summary",
        "LayerCompliance",
        "TextStyleFonts",
        "TextHeights",
        "PurgeUnusedResources",
        "EquipmentTags",
        "FolderDeliverables",
        "DrawingsList",
        "ConversionLog",
        "Errors",
        "CodeCompliance",
    }
    assert expected_sheets.issubset(set(wb.sheetnames))

    layer_ws = wb["LayerCompliance"]
    header = [c.value for c in layer_ws[1]]
    assert header == ["Severity", "File", "Message", "Details", "Occurrences"]
    assert layer_ws.cell(row=2, column=1).value == "FAIL"
