"""End-to-end CLI test using native .dxf fixtures -- no ODA/AutoCAD needed
since there are no .dwg files involved."""
from __future__ import annotations

from pathlib import Path

import openpyxl
from click.testing import CliRunner
from openpyxl import Workbook

from fncsqaqc.cli import cli
from tests import factories


def _write_layers_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Layers"
    ws.append(["Layer Name"])
    ws.append(["M-DUCT"])
    styles = wb.create_sheet("TextStyles")
    styles.append(["Style Name", "Expected Font", "Expected Height"])
    styles.append(["MEP-DUCT-TAG", "simplex.shx", 2.5])
    wb.save(path)


def _write_deliverables_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Deliverables"
    ws.append(["Category", "Item", "Expected Pattern", "Required (Y/N)"])
    ws.append(["Drawings", "HVAC Layout", "*.dxf", "Y"])
    wb.save(path)


def _write_drawings_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Drawings"
    ws.append(["Drawing No.", "Title", "Expected Filename Pattern", "Required (Y/N)"])
    ws.append(["M-101", "HVAC Layout", "M-101.dxf", "Y"])
    wb.save(path)


def test_check_command_end_to_end(tmp_path):
    project_folder = tmp_path / "project"
    project_folder.mkdir()
    doc = factories.make_basic_doc()
    doc.saveas(project_folder / "M-101.dxf")

    layers_xlsx = tmp_path / "layers.xlsx"
    deliverables_xlsx = tmp_path / "deliverables.xlsx"
    drawings_xlsx = tmp_path / "drawings.xlsx"
    _write_layers_workbook(layers_xlsx)
    _write_deliverables_workbook(deliverables_xlsx)
    _write_drawings_workbook(drawings_xlsx)

    output_path = tmp_path / "report.xlsx"

    runner = CliRunner()
    result = runner.invoke(
        cli,
        [
            "check",
            str(project_folder),
            "--layers-excel",
            str(layers_xlsx),
            "--deliverables-excel",
            str(deliverables_xlsx),
            "--drawings-excel",
            str(drawings_xlsx),
            "--output",
            str(output_path),
            "--fail-on",
            "none",
        ],
    )

    assert result.exit_code == 0, result.output
    assert output_path.is_file()
    wb = openpyxl.load_workbook(output_path)
    assert "Summary" in wb.sheetnames
