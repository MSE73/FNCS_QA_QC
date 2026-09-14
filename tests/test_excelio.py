import pytest
from openpyxl import Workbook

from fncsqaqc.excelio.common import ExcelSchemaError
from fncsqaqc.excelio.deliverables_list import load_deliverables
from fncsqaqc.excelio.layer_list import load_layer_standard
from fncsqaqc.excelio.sizing_summary import load_velocity_sizing
from fncsqaqc.excelio.velocity_limits import load_velocity_limits


def test_load_layer_standard_valid(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Layers"
    ws.append(["Layer Name", "Discipline", "Description", "Active (Y/N)"])
    ws.append(["M-DUCT", "Mechanical", "Ductwork", "Y"])
    styles_ws = wb.create_sheet("TextStyles")
    styles_ws.append(["Style Name", "System/Discipline", "Expected Font", "Expected Height"])
    styles_ws.append(["MEP-DUCT-TAG", "Mechanical", "simplex.shx", 2.5])
    path = tmp_path / "layers.xlsx"
    wb.save(path)

    standard = load_layer_standard(path)
    assert standard.active_layer_names() == {"M-DUCT"}
    assert standard.style_by_name()["MEP-DUCT-TAG"].expected_font == "simplex.shx"


def test_load_layer_standard_missing_column_raises_specific_error(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Layers"
    ws.append(["Layer Name Typo", "Discipline"])  # missing "Layer Name"
    ws.append(["M-DUCT", "Mechanical"])
    wb.create_sheet("TextStyles").append(["Style Name"])
    path = tmp_path / "bad_layers.xlsx"
    wb.save(path)

    with pytest.raises(ExcelSchemaError) as exc_info:
        load_layer_standard(path)
    assert "Layer Name" in str(exc_info.value)
    assert "Layers" in str(exc_info.value)


def test_load_deliverables_missing_sheet_raises(tmp_path):
    wb = Workbook()
    wb.active.title = "WrongSheetName"
    path = tmp_path / "deliverables.xlsx"
    wb.save(path)

    with pytest.raises(ExcelSchemaError) as exc_info:
        load_deliverables(path)
    assert "sheet not found" in str(exc_info.value)


def test_load_deliverables_valid(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Deliverables"
    ws.append(["Category", "Item", "Expected Pattern", "Required (Y/N)"])
    ws.append(["Calcs", "HVAC Load", "Calculations/*.pdf", "Y"])
    path = tmp_path / "deliverables.xlsx"
    wb.save(path)

    items = load_deliverables(path)
    assert len(items) == 1
    assert items[0].category == "Calcs"
    assert items[0].required is True


def test_load_velocity_limits_valid(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "VelocityLimits"
    ws.append(["Type (Duct/Pipe)", "System", "Min Velocity (m/s)", "Max Velocity (m/s)", "Code Basis", "Active (Y/N)"])
    ws.append(["Duct", "Supply", None, 8.0, "ASHRAE", "Y"])
    ws.append(["Pipe", "Fuel Gas", None, None, "NFPA 54", "N"])
    path = tmp_path / "velocity_limits.xlsx"
    wb.save(path)

    limits = load_velocity_limits(path)
    assert len(limits) == 2
    supply = next(l for l in limits if l.system == "Supply")
    assert supply.max_velocity == 8.0
    assert supply.active is True
    gas = next(l for l in limits if l.system == "Fuel Gas")
    assert gas.active is False


def test_load_velocity_sizing_valid(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Velocity"
    ws.append(["Tag", "Type (Duct/Pipe)", "System", "Design Flow (L/s)", "Installed Size", "Notes"])
    ws.append(["SD-01", "Duct", "Supply", 850, "600x300", "Main duct"])
    path = tmp_path / "sizing_summary.xlsx"
    wb.save(path)

    rows = load_velocity_sizing(path)
    assert len(rows) == 1
    assert rows[0].tag == "SD-01"
    assert rows[0].design_flow_ls == 850.0


def test_load_velocity_sizing_missing_column_raises(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Velocity"
    ws.append(["Tag", "Type (Duct/Pipe)", "System"])  # missing Design Flow / Installed Size
    path = tmp_path / "bad_sizing.xlsx"
    wb.save(path)

    with pytest.raises(ExcelSchemaError):
        load_velocity_sizing(path)
