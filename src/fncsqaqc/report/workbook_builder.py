"""Builds the Excel QA/QC report from a RunResult."""
from __future__ import annotations

import datetime as dt
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from fncsqaqc import __version__
from fncsqaqc.models import RunResult, Severity

_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_SEVERITY_FILL = {
    Severity.FAIL: PatternFill("solid", fgColor="F8D7DA"),
    Severity.WARN: PatternFill("solid", fgColor="FFF3CD"),
    Severity.INFO: PatternFill("solid", fgColor="D1ECF1"),
}

_SHEET_FOR_CHECK = {
    "layer_names": "LayerCompliance",
    "text_style_fonts": "TextStyleFonts",
    "text_heights": "TextHeights",
    "purge": "PurgeUnusedResources",
    "equipment_tags": "EquipmentTags",
    "folder_deliverables": "FolderDeliverables",
    "drawings_list": "DrawingsList",
    "revit_deferred": "Errors",
}

_RESULT_COLUMNS = ["Severity", "File", "Message", "Details", "Occurrences"]


def _write_header(ws: Worksheet, columns: list[str]) -> None:
    for col_idx, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    ws.freeze_panes = "A2"


def _autosize(ws: Worksheet, columns: list[str], max_width: int = 80) -> None:
    for col_idx, name in enumerate(columns, start=1):
        letter = ws.cell(row=1, column=col_idx).column_letter
        width = max(len(name) + 2, 12)
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=2):
            for cell in row:
                if cell.value is not None:
                    width = max(width, min(len(str(cell.value)) + 2, max_width))
        ws.column_dimensions[letter].width = width


def _write_results_sheet(wb: Workbook, sheet_name: str, results: list) -> None:
    ws = wb.create_sheet(sheet_name)
    _write_header(ws, _RESULT_COLUMNS)
    for row_idx, r in enumerate(results, start=2):
        ws.cell(row=row_idx, column=1, value=r.severity.value)
        ws.cell(row=row_idx, column=2, value=r.file)
        ws.cell(row=row_idx, column=3, value=r.message)
        ws.cell(row=row_idx, column=4, value=r.details)
        ws.cell(row=row_idx, column=5, value=r.occurrences)
        fill = _SEVERITY_FILL.get(r.severity)
        if fill:
            for col in range(1, len(_RESULT_COLUMNS) + 1):
                ws.cell(row=row_idx, column=col).fill = fill
    if not results:
        ws.cell(row=2, column=1, value="No findings.")
    _autosize(ws, _RESULT_COLUMNS)


def _write_summary_sheet(wb: Workbook, run: RunResult) -> None:
    ws = wb.create_sheet("Summary", 0)
    ws.cell(row=1, column=1, value="FNCS QA/QC Report").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=f"fncsqaqc v{__version__}")
    ws.cell(row=3, column=1, value=f"Folder: {run.folder}")
    ws.cell(row=4, column=1, value=f"Run at: {dt.datetime.now().isoformat(timespec='seconds')}")

    counts = run.counts()
    row = 6
    ws.cell(row=row, column=1, value="Overall counts").font = Font(bold=True)
    row += 1
    for sev in (Severity.FAIL, Severity.WARN, Severity.INFO):
        ws.cell(row=row, column=1, value=sev.value)
        ws.cell(row=row, column=2, value=counts[sev])
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="By check").font = Font(bold=True)
    row += 1
    ws.cell(row=row, column=1, value="Check")
    ws.cell(row=row, column=2, value="FAIL")
    ws.cell(row=row, column=3, value="WARN")
    ws.cell(row=row, column=4, value="INFO")
    for c in range(1, 5):
        ws.cell(row=row, column=c).font = Font(bold=True)
    row += 1

    by_check: dict[str, Counter] = defaultdict(Counter)
    for r in run.results:
        by_check[r.check_id][r.severity] += 1
    for check_id in sorted(by_check):
        c = by_check[check_id]
        ws.cell(row=row, column=1, value=check_id)
        ws.cell(row=row, column=2, value=c[Severity.FAIL])
        ws.cell(row=row, column=3, value=c[Severity.WARN])
        ws.cell(row=row, column=4, value=c[Severity.INFO])
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Top failing files").font = Font(bold=True)
    row += 1
    by_file: Counter = Counter()
    for r in run.results:
        if r.severity == Severity.FAIL:
            by_file[r.file] += 1
    for file, count in by_file.most_common(15):
        ws.cell(row=row, column=1, value=file)
        ws.cell(row=row, column=2, value=count)
        row += 1

    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 14


def _write_conversion_log(wb: Workbook, run: RunResult) -> None:
    ws = wb.create_sheet("ConversionLog")
    columns = ["File", "Status", "Detail", "Duration (s)"]
    _write_header(ws, columns)
    for row_idx, entry in enumerate(run.conversion_log, start=2):
        ws.cell(row=row_idx, column=1, value=entry.relative_path)
        ws.cell(row=row_idx, column=2, value=entry.status)
        ws.cell(row=row_idx, column=3, value=entry.detail)
        ws.cell(row=row_idx, column=4, value=round(entry.duration_s, 2))
        if entry.status == "failed":
            for col in range(1, 5):
                ws.cell(row=row_idx, column=col).fill = _SEVERITY_FILL[Severity.FAIL]
    _autosize(ws, columns)


def _write_code_compliance_placeholder(wb: Workbook) -> None:
    ws = wb.create_sheet("CodeCompliance")
    ws.cell(row=1, column=1, value="Reserved for Phase 2 (Jordanian/SBC code compliance checks).").font = Font(
        italic=True
    )


def build(run: RunResult, output_path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)  # placeholder default sheet

    _write_summary_sheet(wb, run)

    by_sheet: dict[str, list] = defaultdict(list)
    for r in run.results:
        by_sheet[_SHEET_FOR_CHECK.get(r.check_id, "Other")].append(r)

    for sheet_name in [
        "LayerCompliance",
        "TextStyleFonts",
        "TextHeights",
        "PurgeUnusedResources",
        "EquipmentTags",
        "FolderDeliverables",
        "DrawingsList",
    ]:
        _write_results_sheet(wb, sheet_name, by_sheet.pop(sheet_name, []))

    for sheet_name, results in by_sheet.items():
        _write_results_sheet(wb, sheet_name, results)

    _write_conversion_log(wb, run)
    _write_results_sheet(wb, "Errors", run.errors)
    _write_code_compliance_placeholder(wb)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
