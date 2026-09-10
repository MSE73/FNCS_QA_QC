"""Shared Excel-reading helpers. Header matching is case-insensitive and
trimmed, and errors name the exact sheet/column at fault -- these workbooks
are hand-maintained by non-programmers, so a typo'd header is the single
most likely real-world failure and deserves a specific, actionable message.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet


class ExcelSchemaError(ValueError):
    def __init__(self, workbook: Path, sheet: str, message: str):
        self.workbook = workbook
        self.sheet = sheet
        super().__init__(f"{workbook.name} [sheet '{sheet}']: {message}")


def _norm(value) -> str:
    return str(value).strip().lower() if value is not None else ""


def load_sheet(workbook_path: Path, sheet_name: str) -> Worksheet:
    if not workbook_path.is_file():
        raise ExcelSchemaError(workbook_path, sheet_name, f"file not found: {workbook_path}")
    wb = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    if sheet_name not in wb.sheetnames:
        raise ExcelSchemaError(
            workbook_path,
            sheet_name,
            f"sheet not found. Available sheets: {', '.join(wb.sheetnames)}",
        )
    return wb[sheet_name]


def read_rows(
    workbook_path: Path, sheet_name: str, required_columns: list[str], optional_columns: list[str] | None = None
) -> list[dict[str, object]]:
    """Reads a sheet's header row, matches it (case-insensitive/trimmed)
    against required/optional columns, and returns one dict per data row."""
    ws = load_sheet(workbook_path, sheet_name)
    optional_columns = optional_columns or []

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        raise ExcelSchemaError(workbook_path, sheet_name, "sheet is empty, expected a header row")

    header_index = {_norm(h): i for i, h in enumerate(header) if h is not None}

    col_index: dict[str, int] = {}
    missing: list[str] = []
    for col in required_columns:
        idx = header_index.get(_norm(col))
        if idx is None:
            missing.append(col)
        else:
            col_index[col] = idx
    if missing:
        raise ExcelSchemaError(
            workbook_path,
            sheet_name,
            f"missing required column(s): {', '.join(missing)}. "
            f"Found headers: {', '.join(str(h) for h in header if h is not None)}",
        )
    for col in optional_columns:
        idx = header_index.get(_norm(col))
        if idx is not None:
            col_index[col] = idx

    rows: list[dict[str, object]] = []
    for raw_row in rows_iter:
        if raw_row is None or all(v is None for v in raw_row):
            continue
        row = {col: (raw_row[idx] if idx < len(raw_row) else None) for col, idx in col_index.items()}
        rows.append(row)

    return rows


def as_bool(value, default: bool = True) -> bool:
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    text = str(value).strip().upper()
    if text in ("Y", "YES", "TRUE", "1"):
        return True
    if text in ("N", "NO", "FALSE", "0"):
        return False
    return default


def as_str(value) -> str:
    return "" if value is None else str(value).strip()


def as_float(value) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None
